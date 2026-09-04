"""Deterministic, local text extraction for project documents.

The project chat never sends whole binary files to an LLM. Files are parsed
locally, converted to bounded text chunks, and the chunks are cited by file
name plus page/sheet/slide when retrieved.
"""
from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree as ET


MAX_DOCUMENT_BYTES = 100 * 1024 * 1024
MAX_EXTRACTED_CHARS = 2_000_000
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
CHUNK_SIZE = 2800
CHUNK_OVERLAP = 250
SUPPORTED_SUFFIXES = {
    ".txt", ".md", ".csv", ".log", ".pdf", ".docx", ".odt", ".xlsx", ".pptx",
}


class DocumentExtractionError(ValueError):
    pass


@dataclass(frozen=True)
class ExtractedPart:
    text: str
    locator: str


def _clean(text: str) -> str:
    text = re.sub(r"[ \t]+", " ", text or "")
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _text_chunks(parts: list[ExtractedPart]) -> list[ExtractedPart]:
    chunks: list[ExtractedPart] = []
    for part in parts:
        text = _clean(part.text)
        if not text:
            continue
        start = 0
        while start < len(text):
            end = min(len(text), start + CHUNK_SIZE)
            if end < len(text):
                boundary = text.rfind("\n", start + CHUNK_SIZE // 2, end)
                if boundary > start:
                    end = boundary
            chunk = text[start:end].strip()
            if chunk:
                chunks.append(ExtractedPart(chunk, part.locator))
            if end >= len(text):
                break
            start = max(start + 1, end - CHUNK_OVERLAP)
    return chunks


def _plain(path: Path) -> list[ExtractedPart]:
    raw = path.read_bytes()
    try:
        text = raw.decode("utf-8")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    return [ExtractedPart(text, "Textdatei")]


def _pdf(path: Path) -> list[ExtractedPart]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - packaging guard
        raise DocumentExtractionError("PDF-Unterstützung ist nicht installiert.") from exc
    try:
        reader = PdfReader(str(path))
        return [ExtractedPart(page.extract_text() or "", f"Seite {i + 1}")
                for i, page in enumerate(reader.pages)]
    except Exception as exc:
        raise DocumentExtractionError(f"PDF konnte nicht gelesen werden: {exc}") from exc


def _docx(path: Path) -> list[ExtractedPart]:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - packaging guard
        raise DocumentExtractionError("DOCX-Unterstützung ist nicht installiert.") from exc
    try:
        document = Document(str(path))
        lines = [p.text for p in document.paragraphs]
        for table in document.tables:
            lines.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
        return [ExtractedPart("\n".join(lines), "Dokument")]
    except Exception as exc:
        raise DocumentExtractionError(f"DOCX konnte nicht gelesen werden: {exc}") from exc


def _xlsx(path: Path) -> list[ExtractedPart]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - packaging guard
        raise DocumentExtractionError("Excel-Unterstützung ist nicht installiert.") from exc
    workbook = None
    try:
        workbook = load_workbook(str(path), read_only=True, data_only=True)
        parts: list[ExtractedPart] = []
        for sheet in workbook.worksheets:
            rows: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                values = [str(value) for value in row if value is not None and str(value).strip()]
                if values:
                    rows.append("\t".join(values))
            parts.append(ExtractedPart("\n".join(rows), f"Tabelle {sheet.title}"))
        return parts
    except Exception as exc:
        raise DocumentExtractionError(f"Excel-Datei konnte nicht gelesen werden: {exc}") from exc
    finally:
        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                pass


def _pptx(path: Path) -> list[ExtractedPart]:
    try:
        from pptx import Presentation
    except ImportError as exc:  # pragma: no cover - packaging guard
        raise DocumentExtractionError("PowerPoint-Unterstützung ist nicht installiert.") from exc
    try:
        presentation = Presentation(str(path))
        parts: list[ExtractedPart] = []
        for index, slide in enumerate(presentation.slides, start=1):
            lines: list[str] = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    lines.append(shape.text)
                if getattr(shape, "has_table", False):
                    lines.extend("\t".join(cell.text for cell in row.cells)
                                 for row in shape.table.rows)
            parts.append(ExtractedPart("\n".join(lines), f"Folie {index}"))
        return parts
    except Exception as exc:
        raise DocumentExtractionError(f"PowerPoint-Datei konnte nicht gelesen werden: {exc}") from exc


def _odt(path: Path) -> list[ExtractedPart]:
    try:
        with zipfile.ZipFile(path) as archive:
            _validate_archive(archive)
            xml = archive.read("content.xml")
        root = ET.fromstring(xml)
        texts = ["".join(node.itertext()) for node in root.iter()
                 if node.tag.endswith("}p") or node.tag.endswith("}h")]
        return [ExtractedPart("\n".join(texts), "Dokument")]
    except Exception as exc:
        raise DocumentExtractionError(f"ODT-Datei konnte nicht gelesen werden: {exc}") from exc


def extract_document(path: Path, filename: str | None = None) -> list[ExtractedPart]:
    """Extract and chunk a supported document, refusing oversized binaries."""
    try:
        path = path.resolve()
        stat = path.stat()
    except (OSError, RuntimeError) as exc:
        raise DocumentExtractionError("Datei konnte nicht geöffnet werden.") from exc
    if not path.is_file():
        raise DocumentExtractionError("Datei nicht gefunden.")
    if stat.st_size > MAX_DOCUMENT_BYTES:
        raise DocumentExtractionError("Datei ist für die lokale Texterkennung zu groß (maximal 100 MiB).")
    suffix = Path(filename or path.name).suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise DocumentExtractionError(f"Dateityp {suffix or 'ohne Endung'} wird nicht unterstützt.")
    if suffix in {".pdf", ".docx", ".odt", ".xlsx", ".pptx"}:
        # OOXML/ODF are ZIP containers; validating their members up front also
        # protects third-party parsers from decompression bombs.
        try:
            with zipfile.ZipFile(path) as archive:
                _validate_archive(archive)
        except zipfile.BadZipFile as exc:
            # PDFs are not ZIP files and are handled by pypdf below.
            if suffix != ".pdf":
                raise DocumentExtractionError("Das Dokument ist kein gültiges Archiv.") from exc
        except (OSError, RuntimeError) as exc:
            raise DocumentExtractionError("Das Dokument konnte nicht geprüft werden.") from exc
    if suffix in {".txt", ".md", ".csv", ".log"}:
        parts = _plain(path)
    elif suffix == ".pdf":
        parts = _pdf(path)
    elif suffix == ".docx":
        parts = _docx(path)
    elif suffix == ".xlsx":
        parts = _xlsx(path)
    elif suffix == ".pptx":
        parts = _pptx(path)
    else:
        parts = _odt(path)
    if sum(len(part.text) for part in parts) > MAX_EXTRACTED_CHARS:
        raise DocumentExtractionError("Die Datei enthält zu viel Text für die lokale Indexierung (maximal 2 Millionen Zeichen).")
    chunks = _text_chunks(parts)
    if not chunks:
        raise DocumentExtractionError("Es wurde kein auslesbarer Text gefunden.")
    return chunks


def _validate_archive(archive: zipfile.ZipFile) -> None:
    """Reject zip bombs and unsafe archive metadata before parser libraries run."""
    total = 0
    for info in archive.infolist():
        # Project documents are parsed, never extracted to disk.  Still reject
        # path traversal and extreme compression ratios to bound parser work.
        name = info.filename.replace("\\", "/")
        if name.startswith("/") or ".." in Path(name).parts:
            raise DocumentExtractionError("Das Dokument enthält einen unsicheren Archivpfad.")
        if info.file_size < 0 or info.file_size > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
            raise DocumentExtractionError("Ein Archivbestandteil ist zu groß.")
        total += info.file_size
        if total > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
            raise DocumentExtractionError("Das Dokument ist nach Dekomprimierung zu groß.")
        if info.compress_size and info.file_size > info.compress_size * 200:
            raise DocumentExtractionError("Verdächtig stark komprimiertes Archiv (möglicher Zip-Bomb).")
