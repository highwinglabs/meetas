"""Meeting export (Phase 2 + Phase 7 selection).

Formats:
* ``markdown`` / ``txt`` / ``json`` / ``html`` -- always available, no deps.
* ``pdf`` / ``docx`` -- *optional*: used only if the local tooling is present
  (``pandoc`` for pdf, ``python-docx`` for docx). Otherwise the export degrades
  cleanly to a close text format (pdf->html, docx->markdown) and reports a
  ``note`` + ``fallback`` so nothing is silently lost and nothing is downloaded.

Every exported statement is traceable: each transcript line carries a stable
segment id and time range, and a "Quellenverweise" section maps each segment id
to its exact time range and source audio. Optional fields render as
"nicht angegeben" (never invented). Exports can include the (validated) LLM
analysis and be filtered to selected sections.
"""
from __future__ import annotations

import html
import json
import os
import shutil
import subprocess
from datetime import datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

from sqlalchemy import select

from core.analysis import schema as analysis_schema
from core.analysis.schema import NOT_GIVEN
from core.config import Config, get_config
from core.logging_setup import get_logger
from core.store.db import session_scope
from core.store.models import Analysis, Meeting, MeetingTag, Recording, TranscriptSegment

log = get_logger("ma.export")

# Always-available text formats + optional binary formats.
FORMATS = ("markdown", "txt", "json", "html", "pdf", "docx")
_TEXT_SUFFIX = {"markdown": ".md", "txt": ".txt", "json": ".json",
                "html": ".html", "pdf": ".pdf", "docx": ".docx"}
_FALLBACK_NOTE = {
    "pdf": "PDF-Tooling (pandoc) nicht vorhanden -- statt PDF wurde HTML erstellt.",
    "docx": "DOCX-Tooling (python-docx) nicht vorhanden -- statt DOCX wurde Markdown erstellt.",
}


def _fmt_ts(seconds: float | None) -> str:
    if seconds is None:
        return "--:--"
    s = max(0, int(round(seconds)))
    return f"{s // 60:02d}:{s % 60:02d}"


def _fmt_dt(dt) -> str:
    if not dt:
        return NOT_GIVEN
    aware = dt if dt.tzinfo is not None else dt.replace(tzinfo=timezone.utc)
    return aware.astimezone(ZoneInfo("Europe/Berlin")).strftime("%Y-%m-%d %H:%M:%S %Z")


def _load(meeting_id: str, include_analysis: bool = True) -> dict:
    with session_scope() as s:
        m = s.get(Meeting, meeting_id)
        if m is None or m.deleted_at is not None:
            raise KeyError(meeting_id)
        rec = s.scalar(select(Recording).where(Recording.meeting_id == meeting_id))
        segs = s.scalars(
            select(TranscriptSegment).where(TranscriptSegment.meeting_id == meeting_id)
            .order_by(TranscriptSegment.start_s)).all()
        tags = [t.tag for t in s.scalars(select(MeetingTag).where(
            MeetingTag.meeting_id == meeting_id).order_by(MeetingTag.tag)).all()]
        data = {
            "meeting": {
                "id": m.id, "title": m.title, "status": m.status, "lang": m.lang,
                "start_at": m.start_at, "end_at": m.end_at, "duration_s": m.duration_s,
            },
            "recording": {
                "source": rec.source, "device": rec.device,
                "sample_rate": rec.sample_rate, "original_path": rec.original_path,
                "status": rec.status,
            } if rec else None,
            "tags": tags,
            "segments": [
                {"id": x.id, "start_s": x.start_s, "end_s": x.end_s,
                 "text": x.text, "speaker_id": x.speaker_id,
                 "language": x.language, "confidence": x.confidence,
                 "audio_ref": x.audio_ref} for x in segs
            ],
        }
        if include_analysis:
            a = s.scalar(select(Analysis).where(
                Analysis.meeting_id == meeting_id, Analysis.kind == "summary"))
            data["analysis"] = analysis_schema.extract_json(a.content) if (a and a.content) else None
        return data


def _speaker(x: dict) -> str:
    return x["speaker_id"] or NOT_GIVEN


def _header_block(data: dict) -> list[str]:
    m, rec = data["meeting"], data["recording"]
    lines = [
        f"- **Status:** {m['status']}",
        f"- **Zeitraum:** {_fmt_dt(m['start_at'])} → {_fmt_dt(m['end_at'])}"
        + (f" ({round(m['duration_s'],1)} s)" if m['duration_s'] else ""),
        f"- **Sprache:** {m['lang'] or NOT_GIVEN}",
        f"- **Audio-Quelle:** {(rec or {}).get('original_path') or NOT_GIVEN}",
        f"- **Segmente:** {len(data['segments'])}",
    ]
    if data.get("tags"):
        lines.append(f"- **Tags:** {', '.join(data['tags'])}")
    return lines


def render_markdown(data: dict, include_transcript: bool = True,
                    include_analysis: bool = True,
                    section_keys: list | None = None) -> str:
    m = data["meeting"]
    lines = [f"# {m['title']}", ""] + _header_block(data) + [""]
    if include_analysis:
        am = analysis_schema.render_markdown(data.get("analysis"), section_keys)
        if am:
            lines += ["## Analyse (lokal, mit Quellen)", "", am.rstrip(), ""]
    if include_transcript:
        lines += ["## Transkript", "",
                  "| Zeit | Sprecher | Text |", "|---|---|---|"]
        for x in data["segments"]:
            speaker = _speaker(x).replace("|", "\\|")
            text = (x["text"] or "").replace("|", "\\|")
            lines.append(f"| {_fmt_ts(x['start_s'])} | {speaker} | {text} `[seg:{x['id']}]` |")
    lines += ["", "## Quellenverweise", ""]
    for x in data["segments"]:
        audio = x["audio_ref"] or NOT_GIVEN
        lines.append(f"- `[seg:{x['id']}]` {_fmt_ts(x['start_s'])}–{_fmt_ts(x['end_s'])} · "
                     f"{audio} · „{x['text']}“")
    return "\n".join(lines) + "\n"


def render_txt(data: dict, include_transcript: bool = True,
               include_analysis: bool = True,
               section_keys: list | None = None) -> str:
    m = data["meeting"]
    lines = [m["title"], "=" * max(4, len(m["title"]))]
    for h in _header_block(data):
        lines.append(h.replace("**", ""))
    if include_analysis:
        am = analysis_schema.render_markdown(data.get("analysis"), section_keys)
        if am:
            lines += ["", "Analyse (lokal, mit Quellen)", "-" * 29]
            lines += [ln.replace("## ", "").replace("##", "")
                      for ln in am.splitlines()]
    if include_transcript:
        lines += ["", "Transkript", "-" * 8]
        for x in data["segments"]:
            lines.append(f"[{_fmt_ts(x['start_s'])}] {_speaker(x)}: {x['text']}  (seg:{x['id']})")
    lines += ["", "Quellenverweise", "-" * 15]
    for x in data["segments"]:
        audio = x["audio_ref"] or NOT_GIVEN
        lines.append(f"seg:{x['id']}  {_fmt_ts(x['start_s'])}-{_fmt_ts(x['end_s'])}  {audio}")
    return "\n".join(lines) + "\n"


def render_json(data: dict, include_transcript: bool = True,
                include_analysis: bool = True,
                section_keys: list | None = None) -> str:
    payload = {
        "export": {"format": "json",
                   "generated_at": datetime.now(timezone.utc).isoformat()},
        "meeting": {**data["meeting"],
                    "start_at": _fmt_dt(data["meeting"]["start_at"]),
                    "end_at": _fmt_dt(data["meeting"]["end_at"])},
        "recording": data["recording"],
        "tags": data.get("tags", []),
    }
    if include_analysis:
        keys = section_keys or list(analysis_schema.SECTION_KEYS)
        payload["analysis"] = {k: data.get("analysis", {}).get(k, [])
                               for k in keys} if data.get("analysis") else None
    if include_transcript:
        payload["segments"] = data["segments"]
    return json.dumps(payload, ensure_ascii=False, indent=2, default=str)


def render_html(data: dict, include_transcript: bool = True,
                include_analysis: bool = True,
                section_keys: list | None = None) -> str:
    m = data["meeting"]
    esc = html.escape
    parts = [
        "<!doctype html><html lang='de'><head><meta charset='utf-8'>",
        f"<title>{esc(m['title'])}</title>",
        "<style>body{font-family:sans-serif;margin:2rem;max-width:900px;line-height:1.4}"
        "table{border-collapse:collapse;width:100%}td,th{border:1px solid #ccc;"
        "padding:.35rem .5rem;vertical-align:top}th{background:#f0f0f0}"
        ".meta li{margin:.1rem 0}</style></head><body>",
        f"<h1>{esc(m['title'])}</h1>",
        "<ul class='meta'>",
    ]
    for h in _header_block(data):
        txt = h.replace("**", "")
        parts.append(f"<li>{esc(txt)}</li>")
    parts.append("</ul>")
    if include_analysis:
        am = analysis_schema.render_markdown(data.get("analysis"), section_keys)
        if am:
            parts.append("<h2>Analyse (lokal, mit Quellen)</h2>")
            parts.append("<pre style='white-space:pre-wrap'>" + esc(am) + "</pre>")
    if include_transcript:
        parts.append("<h2>Transkript</h2><table><tr><th>Zeit</th><th>Sprecher</th>"
                     "<th>Text</th></tr>")
        for x in data["segments"]:
            parts.append(
                f"<tr><td>{_fmt_ts(x['start_s'])}</td><td>{esc(_speaker(x))}</td>"
                f"<td>{esc(x['text'] or '')} <code>[seg:{x['id']}]</code></td></tr>")
        parts.append("</table>")
    parts.append("<h2>Quellenverweise</h2><ul>")
    for x in data["segments"]:
        audio = x["audio_ref"] or NOT_GIVEN
        parts.append(f"<li><code>[seg:{x['id']}]</code> "
                     f"{_fmt_ts(x['start_s'])}–{_fmt_ts(x['end_s'])} · {esc(audio)}</li>")
    parts += ["</ul>", "</body></html>"]
    return "\n".join(parts) + "\n"


def _pandoc_available() -> bool:
    return shutil.which("pandoc") is not None


_PDF_ENGINES = ("pdflatex", "xelatex", "lualatex", "wkhtmltopdf", "weasyprint")


def _pdf_engine() -> str | None:
    """First PDF engine installed on this system (pandoc 3.x does not
    fall back automatically, so the engine must be passed explicitly)."""
    for name in _PDF_ENGINES:
        if shutil.which(name) is not None:
            return name
    return None


def _build_docx(data: dict, path: Path, md: str,
                include_analysis: bool = True,
                include_transcript: bool = True,
                section_keys: list | None = None) -> bool:
    """Build a .docx via python-docx if available. Returns False if not installed."""
    try:
        import docx  # python-docx
    except Exception:
        return False
    d = docx.Document()
    m = data["meeting"]
    d.add_heading(m["title"], 0)
    for h in _header_block(data):
        d.add_paragraph(h.replace("**", ""))
    # Keep the binary export in lockstep with the text exporters: callers may
    # deliberately omit the analysis or transcript, and section filtering must
    # apply to DOCX as well.
    if include_analysis and data.get("analysis"):
        d.add_heading("Analyse", level=1)
        am = analysis_schema.render_markdown(data.get("analysis"), section_keys)
        if am:
            for ln in am.splitlines():
                if ln.startswith("## "):
                    d.add_heading(ln[3:], level=2)
                else:
                    d.add_paragraph(ln)
    if include_transcript:
        d.add_heading("Transkript", level=1)
        for x in data["segments"]:
            d.add_paragraph(f"[{_fmt_ts(x['start_s'])}] {_speaker(x)}: {x['text']}  (seg:{x['id']})")
        d.add_heading("Quellenverweise", level=1)
        for x in data["segments"]:
            audio = x["audio_ref"] or NOT_GIVEN
            d.add_paragraph(f"seg:{x['id']}  {_fmt_ts(x['start_s'])}-{_fmt_ts(x['end_s'])}  {audio}")
    d.save(str(path))
    return True


def build_export(meeting_id: str, fmt: str = "markdown", config: Config | None = None,
                 dest_dir: Path | None = None, include_analysis: bool = True,
                 include_transcript: bool = True,
                 section_keys: list | None = None) -> dict:
    """Export a meeting, honouring format + content selection.

    Returns ``{path, format, suffix, is_binary, note, content}``. Optional
    formats (pdf/docx) fall back cleanly when the local tooling is absent.
    """
    if fmt not in FORMATS:
        raise ValueError(f"unbekanntes Format: {fmt}")
    cfg = config or get_config()
    allowed = cfg.export_formats or list(FORMATS)
    if fmt not in allowed and fmt not in ("pdf", "docx"):
        raise ValueError(f"Format '{fmt}' ist nicht freigeschaltet (erlaubt: {allowed})")
    cfg.ensure_dirs()
    data = _load(meeting_id, include_analysis=include_analysis)
    fallback = None
    note = None

    if fmt in ("markdown", "txt", "json", "html"):
        content = {
            "markdown": render_markdown, "txt": render_txt,
            "json": render_json, "html": render_html,
        }[fmt](data, include_transcript=include_transcript,
               include_analysis=include_analysis, section_keys=section_keys)
        out_bytes, is_binary = content.encode("utf-8"), False
    elif fmt == "pdf":
        engine = _pdf_engine() if _pandoc_available() else None
        if engine is not None:
            md = render_markdown(data, include_transcript=include_transcript,
                                 include_analysis=include_analysis,
                                 section_keys=section_keys)
            try:
                try:
                    out_bytes = subprocess.run(
                        ["pandoc", "-f", "markdown", "-t", "pdf",
                         f"--pdf-engine={engine}"],
                        input=md.encode("utf-8"), capture_output=True, check=True,
                        timeout=300).stdout
                except subprocess.CalledProcessError as exc:
                    raw_detail = exc.stderr or b""
                    detail = (raw_detail.decode("utf-8", "replace")
                              if isinstance(raw_detail, bytes) else str(raw_detail))[:300]
                    raise ValueError(f"PDF-Export fehlgeschlagen: {detail}") from exc
            except subprocess.TimeoutExpired as exc:
                raise ValueError("PDF-Export dauerte zu lange und wurde beendet.") from exc
        else:
            fallback = "html"
            note = (_FALLBACK_NOTE["pdf"] if not _pandoc_available()
                    else "PDF-Engine (pdflatex/wkhtmltopdf/weasyprint) nicht "
                         "gefunden -- statt PDF wurde HTML erstellt.")
            out_bytes = render_html(data, include_transcript=include_transcript,
                                    include_analysis=include_analysis,
                                    section_keys=section_keys).encode("utf-8")
        is_binary = fallback is None
    else:  # docx
        md = render_markdown(data, include_transcript=include_transcript,
                             include_analysis=include_analysis,
                             section_keys=section_keys)
        tmp = Path(cfg.exports_dir) / f"{meeting_id[:8]}_tmp.docx"
        tmp.parent.mkdir(parents=True, exist_ok=True)
        try:
            built_docx = _build_docx(data, tmp, md, include_analysis,
                                     include_transcript, section_keys)
        except Exception as exc:
            tmp.unlink(missing_ok=True)
            raise ValueError("DOCX-Export fehlgeschlagen.") from exc
        if built_docx:
            try:
                tmp.chmod(0o600)
            except OSError:
                pass
            out_bytes, is_binary = tmp.read_bytes(), True
            tmp.unlink(missing_ok=True)
        else:
            tmp.unlink(missing_ok=True)
            fallback = "markdown"
            note = _FALLBACK_NOTE["docx"]
            out_bytes = md.encode("utf-8")
            is_binary = False

    effective_fmt = fallback or fmt
    suffix = _TEXT_SUFFIX[effective_fmt]
    out_dir = Path(dest_dir) if dest_dir else cfg.exports_dir
    out_dir.mkdir(parents=True, exist_ok=True)
    m = data["meeting"]
    safe_title = "".join(c if c.isalnum() or c in " -_" else "_"
                         for c in (m["title"] or "meeting"))[:40].strip() or "meeting"
    out_path = out_dir / f"{m['id'][:8]}_{safe_title}{suffix}"
    # Atomic replacement prevents a crash or interrupted export from leaving a
    # truncated file at the user-visible destination.
    tmp_path = out_path.with_name(out_path.name + ".tmp")
    try:
        with open(tmp_path, "wb") as handle:
            handle.write(out_bytes)
            handle.flush()
            os.fsync(handle.fileno())
        os.replace(tmp_path, out_path)
    except Exception:
        tmp_path.unlink(missing_ok=True)
        raise
    try:
        out_path.chmod(0o600)
    except OSError:
        pass
    log.info("export_done meeting=%s fmt=%s effective=%s path=%s",
             meeting_id[:8], fmt, effective_fmt, out_path)
    return {
        "path": str(out_path), "format": effective_fmt, "suffix": suffix,
        "is_binary": is_binary, "note": note, "fallback": fallback,
        "content": out_bytes.decode("utf-8") if not is_binary else None,
        "size": len(out_bytes),
    }


def export_meeting(meeting_id: str, fmt: str = "markdown", config: Config | None = None,
                   dest_dir: Path | None = None) -> Path:
    """Backward-compatible wrapper returning the written path."""
    return Path(build_export(meeting_id, fmt=fmt, config=config, dest_dir=dest_dir)["path"])
