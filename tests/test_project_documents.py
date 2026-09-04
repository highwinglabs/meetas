"""Local project-document extraction, including spreadsheet and presentation files."""
from __future__ import annotations

from core.project_documents import extract_document


def test_extract_excel_and_powerpoint(tmp_path):
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Plan"
    sheet.append(["Aufgabe", "Verantwortlich"])
    sheet.append(["Angebot prüfen", "Lea"])
    xlsx = tmp_path / "plan.xlsx"
    workbook.save(xlsx)
    workbook.close()

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    box.text = "Projektstatus: Freigabe am Freitag"
    pptx = tmp_path / "status.pptx"
    presentation.save(pptx)

    excel_parts = extract_document(xlsx)
    ppt_parts = extract_document(pptx)
    assert any("Angebot prüfen" in part.text and "Lea" in part.text
               and "Plan" in part.locator for part in excel_parts)
    assert any("Freigabe am Freitag" in part.text and "Folie 1" in part.locator
               for part in ppt_parts)


def test_project_file_is_indexed_and_reindexable(make_service, tmp_path):
    from openpyxl import Workbook
    from core.store.db import session_scope

    svc = make_service()
    project = svc.create_project("Dokumentprojekt")
    workbook = Workbook()
    workbook.active.append(["Budget", "Freigabe nächste Woche"])
    source = tmp_path / "budget.xlsx"
    workbook.save(source)
    workbook.close()

    result = svc.import_upload(source.name, source.read_bytes(), project_id=project["id"],
                               content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert result["project_id"] == project["id"]
    detail = svc.project_detail(project["id"])
    assert detail["files"][0]["extraction_status"] == "ready"
    assert detail["files"][0]["chunks"] >= 1
    with session_scope() as session:
        hits = svc._project_document_hits(session, project["id"], "Budget Freigabe")
    assert hits and hits[0]["source_kind"] == "document"
    assert hits[0]["file_name"] == "budget.xlsx"
    assert hits[0]["locator"] == "Tabelle Sheet"
    assert svc.reindex_project_files(project["id"])["ready"] == 1


def test_project_file_trash_restore_and_permanent_delete(make_service, tmp_path):
    svc = make_service()
    project = svc.create_project("Papierkorb")
    source = tmp_path / "notiz.txt"
    source.write_text("Wichtige Projektentscheidung", encoding="utf-8")
    result = svc.import_upload(source.name, source.read_bytes(), project_id=project["id"])
    file_id = result["id"]

    assert svc.trash_project_file(file_id)["id"] == file_id
    assert svc.project_detail(project["id"])["files"] == []
    trash = svc.list_trash()
    assert any(row["id"] == file_id for row in trash["files"])
    assert svc.restore_project_file(file_id)["restored"] is True
    assert svc.project_detail(project["id"])["files"][0]["id"] == file_id

    svc.trash_project_file(file_id)
    assert svc.permanently_delete_project_file(file_id)["deleted"] is True
    assert svc.project_detail(project["id"])["files"] == []
    assert not any(row["id"] == file_id for row in svc.list_trash()["files"])
